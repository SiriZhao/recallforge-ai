from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from app.schemas.course_model import SourceRef
from app.schemas.review import ParsedPage
from app.services.file_parser import parse_file
from app.services.structure_analyzer import analyze_document_structure
from app.services.evidence_catalog import catalog_for, resolves


def test_docx_preserves_body_table_order_and_content_hash(tmp_path):
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("Before experiment")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "25 C"
    document.add_paragraph("After experiment")
    document.save(path)
    parsed = parse_file(path)
    assert parsed.raw_text.index("Before") < parsed.raw_text.index("Control") < parsed.raw_text.index("After")
    assert len(parsed.content_hash) == 64
    assert parsed.content_hash == parse_file(path).content_hash


def test_ppt_table_reaches_generation_text_and_title_is_structural(tmp_path):
    path = tmp_path / "synthetic.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Transport experiment"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Temperature"
    table.cell(1, 0).text = "25 C"
    deck.save(path)
    parsed = parse_file(path)
    assert "Temperature" in parsed.raw_text and "25 C" in parsed.raw_text
    assert any(section.title == "Transport experiment" for section in parsed.document_structure.sections)


def test_processing_windows_do_not_claim_semantic_chapters():
    pages = [ParsedPage(page_number=i, text="Body text without headings.", source="text_extract") for i in range(1, 45)]
    structure = analyze_document_structure("long.txt", ".txt", pages)
    assert len(structure.sections) > 1
    assert all(section.basis == "processing_window" for section in structure.sections)


def test_evidence_reference_checks_page_fragment_block_and_offsets(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("A membrane controls transport.")
    spans = catalog_for([parse_file(path)])
    ref = spans[0].source.model_copy(update={"fragment": "membrane", "text_start": 2, "text_end": 10})
    assert resolves(ref, spans)
    for update in ({"page": 99}, {"filename": "invented.txt"}, {"fragment": "fake claim"}, {"block_id": "invented"}, {"text_end": 1000}):
        assert not resolves(ref.model_copy(update=update), spans)
